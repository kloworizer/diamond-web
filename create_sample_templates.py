"""Script to create sample DOCX template files with proper placeholders."""

import os
from docx import Document
from pathlib import Path

# Define templates with their placeholder content
TEMPLATES = {
    'tanda_terima_nasional_internasional': {
        'nama': 'Tanda Terima ILAP Nasional/Internasional',
        'deskripsi': 'Template tanda terima untuk ILAP nasional dan internasional',
        'content': '''TANDA TERIMA DATA ILAP NASIONAL/INTERNASIONAL

Nomor Tanda Terima: {{nomor_tanda_terima}}
Tanggal: {{tanggal_terima_dip}}

Diterima dari: {{diterima_dari}}
Nomor Surat Pengantar: {{nomor_surat_pengantar}}
Tanggal Surat Pengantar: {{tanggal_surat_pengantar}}

ILAP/Data: {{nama_ilap}}
Jenis Data: {{jenis_data}}
Periode Data: {{periode_data}}
Bentuk Data: {{bentuk_data}}
Cara Penyampaian: {{cara_penyampaian}}

PIC P3DE: {{nama_pic_p3de}}

Status: Diterima'''
    },
    'tanda_terima_regional': {
        'nama': 'Tanda Terima ILAP Regional',
        'deskripsi': 'Template tanda terima untuk ILAP regional',
        'content': '''TANDA TERIMA DATA ILAP REGIONAL

Nomor Tanda Terima: {{nomor_tanda_terima}}
Tanggal: {{tanggal_terima_dip}}

Diterima dari: {{diterima_dari}}
Nomor Surat Pengantar: {{nomor_surat_pengantar}}
Tanggal Surat Pengantar: {{tanggal_surat_pengantar}}

Data: {{nama_ilap}}
Jenis Data: {{jenis_data}}
Periode: {{periode_data}}
Bentuk: {{bentuk_data}}
Penyampaian: {{cara_penyampaian}}

Petugas P3DE: {{nama_pic_p3de}}
Tanggal Diterima: {{tanggal_terima_dip}}

---
Tanda Tangan Penerima'''
    },
    'lampiran_tanda_terima_nasional_internasional': {
        'nama': 'Lampiran Tanda Terima ILAP Nasional/Internasional',
        'deskripsi': 'Lampiran rincian tanda terima untuk ILAP nasional dan internasional',
        'content': '''LAMPIRAN TANDA TERIMA DATA ILAP NASIONAL/INTERNASIONAL

Nomor Tanda Terima: {{nomor_tanda_terima}}

Rincian Data:
- ILAP: {{nama_ilap}}
- Jenis Data: {{jenis_data}}
- Periode: {{periode_data}}
- Bentuk: {{bentuk_data}}
- Cara Penyampaian: {{cara_penyampaian}}

Pengirim: {{diterima_dari}}
Penerima: {{nama_pic_p3de}}

Keterangan Lainnya:
{{keterangan_lainnya}}'''
    },
    'lampiran_tanda_terima_regional': {
        'nama': 'Lampiran Tanda Terima ILAP Regional',
        'deskripsi': 'Lampiran rincian tanda terima untuk ILAP regional',
        'content': '''LAMPIRAN TANDA TERIMA DATA ILAP REGIONAL

Nomor Tanda Terima: {{nomor_tanda_terima}}
Tanggal: {{tanggal_terima_dip}}

Data yang Diterima:
Nama: {{nama_ilap}}
Jenis: {{jenis_data}}
Periode: {{periode_data}}
Bentuk: {{bentuk_data}}

Dari: {{diterima_dari}}
Kepada: {{nama_pic_p3de}}

Catatan:
{{keterangan_lainnya}}'''
    },
    'register_penerimaan_data': {
        'nama': 'Register Penerimaan Data',
        'deskripsi': 'Register pencatatan penerimaan data P3DE',
        'content': '''REGISTER PENERIMAAN DATA P3DE

Nomor Tanda Terima: {{nomor_tanda_terima}}
Tanggal Terima: {{tanggal_terima_dip}}

Pengirim: {{diterima_dari}}
Data/ILAP: {{nama_ilap}}
Jenis Data: {{jenis_data}}
Periode: {{periode_data}}

Bentuk: {{bentuk_data}}
Cara: {{cara_penyampaian}}

Penerima: {{nama_pic_p3de}}

Status Diterima: ✓ Lengkap
Catatan: -'''
    },
    'nd_pengantar_pide': {
        'nama': 'ND Pengantar ke PIDE',
        'deskripsi': 'Naskah Dinas pengantar pengiriman data ke PIDE',
        'content': '''NASKAH DINAS PENGANTAR DATA KE PIDE

Nomor: {{nomor_surat_pengantar}}
Tanggal: {{tanggal_surat_pengantar}}

Kepada: Kepala PIDE

Sehubungan dengan data ILAP yang telah diterima:

Nama ILAP: {{nama_ilap}}
Jenis Data: {{jenis_data}}
Periode: {{periode_data}}
Bentuk Data: {{bentuk_data}}

Nomor Tanda Terima: {{nomor_tanda_terima}}
Diterima dari: {{diterima_dari}}

Data tersebut telah selesai diproses dan siap untuk pengiriman ke PIDE.

Demikian untuk menjadi perhatian.

{{nama_pic_p3de}}
Petugas P3DE'''
    },
    'surat_klarifikasi': {
        'nama': 'Surat Klarifikasi',
        'deskripsi': 'Surat klarifikasi data untuk pengirim',
        'content': '''SURAT KLARIFIKASI DATA

Nomor: {{nomor_surat_pengantar}}
Tanggal: {{tanggal_surat_pengantar}}

Kepada: {{diterima_dari}}

Dengan hormat,

Berkenaan dengan pengiriman data {{nama_ilap}} dengan nomor tanda terima {{nomor_tanda_terima}}, 
diperlukan klarifikasi terkait:

Data: {{nama_ilap}}
Periode: {{periode_data}}
Bentuk: {{bentuk_data}}

Mohon dapat dikonfirmasi dan dikirimkan kembali jika diperlukan perbaikan.

Terima kasih atas perhatian dan kerjasamanya.

{{nama_pic_p3de}}
Petugas P3DE'''
    },
    'surat_pkdi_nasional_internasional_lengkap': {
        'nama': 'Surat PKDI ILAP Nasional/Internasional Lengkap',
        'deskripsi': 'Surat Pernyataan Kesesuaian Data ILAP Lengkap',
        'content': '''SURAT PERNYATAAN KESESUAIAN DATA ILAP
NASIONAL/INTERNASIONAL - DATA LENGKAP

Nomor: {{nomor_surat_pengantar}}
Tanggal: {{tanggal_surat_pengantar}}

Dengan ini kami menyatakan bahwa data {{nama_ilap}} yang diterima pada {{tanggal_terima_dip}}:

Jenis Data: {{jenis_data}}
Periode: {{periode_data}}
Bentuk: {{bentuk_data}}
Status: LENGKAP

Semua variabel dan satuan data telah sesuai dengan spesifikasi yang diminta.

Nomor Tanda Terima: {{nomor_tanda_terima}}
Diterima dari: {{diterima_dari}}

{{nama_pic_p3de}}
Petugas P3DE'''
    },
    'surat_pkdi_nasional_internasional_sebagian': {
        'nama': 'Surat PKDI ILAP Nasional/Internasional Lengkap Sebagian',
        'deskripsi': 'Surat Pernyataan Kesesuaian Data ILAP Lengkap Sebagian',
        'content': '''SURAT PERNYATAAN KESESUAIAN DATA ILAP
NASIONAL/INTERNASIONAL - DATA LENGKAP SEBAGIAN

Nomor: {{nomor_surat_pengantar}}
Tanggal: {{tanggal_surat_pengantar}}

Dengan ini kami menyatakan bahwa data {{nama_ilap}} yang diterima pada {{tanggal_terima_dip}}:

Jenis Data: {{jenis_data}}
Periode: {{periode_data}}
Bentuk: {{bentuk_data}}
Status: LENGKAP SEBAGIAN

Sebagian variabel dan satuan data telah sesuai dengan spesifikasi yang diminta.

Nomor Tanda Terima: {{nomor_tanda_terima}}
Diterima dari: {{diterima_dari}}

Keterangan yang tidak lengkap akan dikomunikasikan lebih lanjut.

{{nama_pic_p3de}}
Petugas P3DE'''
    },
    'surat_pkdi_regional_lengkap': {
        'nama': 'Surat PKDI ILAP Regional Lengkap',
        'deskripsi': 'Surat Pernyataan Kesesuaian Data ILAP Regional Lengkap',
        'content': '''SURAT PERNYATAAN KESESUAIAN DATA ILAP REGIONAL
DATA LENGKAP

Nomor: {{nomor_surat_pengantar}}
Tanggal: {{tanggal_surat_pengantar}}

Dengan ini kami menyatakan bahwa data {{nama_ilap}} (Regional):

Jenis Data: {{jenis_data}}
Periode: {{periode_data}}
Bentuk Data: {{bentuk_data}}

Yang diterima pada {{tanggal_terima_dip}} dengan nomor tanda terima {{nomor_tanda_terima}}

Status: LENGKAP

Semua unit dan variabel telah memenuhi standar data yang ditetapkan.

Pengirim Data: {{diterima_dari}}
Petugas Penerima: {{nama_pic_p3de}}

Demikian surat ini dibuat untuk kegiatan administrasi pendataan.'''
    },
    'surat_pkdi_regional_sebagian': {
        'nama': 'Surat PKDI ILAP Regional Lengkap Sebagian',
        'deskripsi': 'Surat Pernyataan Kesesuaian Data ILAP Regional Lengkap Sebagian',
        'content': '''SURAT PERNYATAAN KESESUAIAN DATA ILAP REGIONAL
DATA LENGKAP SEBAGIAN

Nomor: {{nomor_surat_pengantar}}
Tanggal: {{tanggal_surat_pengantar}}

Dengan ini kami menyatakan bahwa data {{nama_ilap}} (Regional):

Jenis Data: {{jenis_data}}
Periode: {{periode_data}}
Bentuk Data: {{bentuk_data}}

Yang diterima pada {{tanggal_terima_dip}} dengan nomor tanda terima {{nomor_tanda_terima}}

Status: LENGKAP SEBAGIAN

Beberapa unit dan variabel masih perlu perbaikan dan akan dikomunikasikan dengan pengirim.

Pengirim Data: {{diterima_dari}}
Petugas Penerima: {{nama_pic_p3de}}

Mohon segera melakukan perbaikan yang diperlukan.'''
    },
}

def create_templates():
    """Create sample DOCX template files."""
    base_dir = Path('d:/diamond-web/diamond_web/media/docx_templates')
    base_dir.mkdir(parents=True, exist_ok=True)
    
    created_files = []
    
    for template_key, template_data in TEMPLATES.items():
        # Create document
        doc = Document()
        doc.add_paragraph(template_data['content'])
        
        # Save file
        filename = f"{template_key}.docx"
        filepath = base_dir / filename
        doc.save(str(filepath))
        
        created_files.append({
            'key': template_key,
            'nama': template_data['nama'],
            'deskripsi': template_data['deskripsi'],
            'filename': filename
        })
        
        print(f"✓ Created: {filename}")
    
    return created_files

if __name__ == '__main__':
    files = create_templates()
    print(f"\n✓ Created {len(files)} template files successfully!")
    for f in files:
        print(f"  - {f['nama']}")
