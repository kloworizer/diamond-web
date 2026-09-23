"""PMDE: Generate ND Pengantar ke PDI for adhoc data."""

from datetime import date, datetime, timedelta
from io import BytesIO

import pytest
from django.core.files.base import ContentFile
from django.urls import reverse
from docx import Document

from diamond_web.constants.tiket_action_types import TiketActionType
from diamond_web.constants.tiket_status import STATUS_PENGENDALIAN_MUTU, STATUS_SELESAI
from diamond_web.models.docx_template import DocxTemplate
from diamond_web.models.tiket_action import TiketAction
from diamond_web.tests.conftest import (
    JenisDataILAPFactory,
    PeriodeJenisDataFactory,
    TiketFactory,
    UserFactory,
)

URL_NAME = 'bulk_nd_pengantar_pdi'
DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def _ago(days):
    return datetime.now() - timedelta(days=days)


def _tiket(nama_tabel_i, selesai_days_ago=(), status_tiket=STATUS_SELESAI):
    """A tiket with one SELESAI action per entry of `selesai_days_ago`.

    It is received long before, so only the Selesai action can put it in
    the one-month window.
    """
    sub = JenisDataILAPFactory(nama_tabel_I=nama_tabel_i)
    periode = PeriodeJenisDataFactory(id_sub_jenis_data_ilap=sub)
    tiket = TiketFactory(id_periode_data=periode, status_tiket=status_tiket, tgl_terima_dip=_ago(200))
    for days in selesai_days_ago:
        TiketAction.objects.create(
            id_tiket=tiket, id_user=UserFactory(), timestamp=_ago(days),
            action=TiketActionType.SELESAI, catatan='selesai',
        )
    return tiket


@pytest.fixture(autouse=True)
def _media_root(settings, tmp_path):
    settings.MEDIA_ROOT = str(tmp_path)


@pytest.fixture
def tikets(db):
    return {
        'adhoc_recent': _tiket('KPDE_ADHOC_BI_PADAN', selesai_days_ago=[3]),
        'adhoc_lowercase': _tiket('kpde_adhoc_lain', selesai_days_ago=[10]),
        # Finished twice: the latest finish is what counts.
        'adhoc_refinished': _tiket('KPDE_ADHOC_ULANG', selesai_days_ago=[50, 5]),
        'adhoc_old': _tiket('KPDE_ADHOC_LAMA', selesai_days_ago=[45]),
        # Reopened after a recent finish: no longer Selesai.
        'adhoc_reopened': _tiket('KPDE_ADHOC_BUKA', selesai_days_ago=[3], status_tiket=STATUS_PENGENDALIAN_MUTU),
        # Status Selesai but no Selesai action in the trail: no finish date.
        'adhoc_no_action': _tiket('KPDE_ADHOC_TANPA_AKSI'),
        'regular_recent': _tiket('KPDE_PBB_P2', selesai_days_ago=[3]),
    }


def _nd_template(paragraph_text):
    doc = Document()
    doc.add_paragraph(paragraph_text)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = '{{row.nomor_tiket}}'
    table.rows[0].cells[1].text = '{{row.nama_tabel_i}}'
    buffer = BytesIO()
    doc.save(buffer)
    template = DocxTemplate(nama_template='ND Pengantar ke PDI', jenis_dokumen='nd_pengantar_pdi', active=True)
    template.file_template.save('nd_pengantar_pdi_test.docx', ContentFile(buffer.getvalue()), save=True)
    return template


@pytest.mark.django_db
class TestAccess:
    def test_pmde_user_allowed(self, client, pmde_user):
        client.force_login(pmde_user)
        assert client.get(reverse(URL_NAME)).status_code == 200

    def test_pmde_admin_allowed(self, client, pmde_admin_user):
        client.force_login(pmde_admin_user)
        assert client.get(reverse(URL_NAME)).status_code == 200

    def test_p3de_user_rejected(self, client, authenticated_user):
        client.force_login(authenticated_user)
        assert client.get(reverse(URL_NAME)).status_code == 302

    def test_anonymous_redirected_to_login(self, client):
        resp = client.get(reverse(URL_NAME))
        assert resp.status_code == 302
        assert 'login' in resp.url


@pytest.mark.django_db
class TestListing:
    def test_default_shows_adhoc_from_last_month(self, client, pmde_user, tikets):
        client.force_login(pmde_user)
        resp = client.get(reverse(URL_NAME))

        shown = {t.pk for t in resp.context['tickets']}
        assert shown == {
            tikets['adhoc_recent'].pk, tikets['adhoc_lowercase'].pk, tikets['adhoc_refinished'].pk,
        }
        today = date.today()
        assert resp.context['selected_tanggal_selesai'] == today.isoformat()
        assert resp.context['selected_tanggal_mulai'] < (today - timedelta(days=27)).isoformat()

    def test_date_range_widens_to_older_data(self, client, pmde_user, tikets):
        client.force_login(pmde_user)
        resp = client.get(reverse(URL_NAME), {
            'tanggal_mulai': (date.today() - timedelta(days=60)).isoformat(),
            'tanggal_selesai': date.today().isoformat(),
        })
        shown = {t.pk for t in resp.context['tickets']}
        assert tikets['adhoc_old'].pk in shown
        assert tikets['regular_recent'].pk not in shown
        assert tikets['adhoc_reopened'].pk not in shown
        assert tikets['adhoc_no_action'].pk not in shown

    def test_row_carries_latest_finish_date(self, client, pmde_user, tikets):
        client.force_login(pmde_user)
        resp = client.get(reverse(URL_NAME))
        row = next(t for t in resp.context['tickets'] if t.pk == tikets['adhoc_refinished'].pk)
        assert row.tgl_selesai.date() == _ago(5).date()
        assert _ago(5).strftime('%d/%m/%Y') in resp.content.decode()

    def test_ilap_filter_and_options(self, client, pmde_user, tikets):
        client.force_login(pmde_user)
        ilap = tikets['adhoc_recent'].id_periode_data.id_sub_jenis_data_ilap.id_ilap
        resp = client.get(reverse(URL_NAME), {'ilap_id': ilap.pk})

        assert [t.pk for t in resp.context['tickets']] == [tikets['adhoc_recent'].pk]
        # Options list every ILAP with adhoc data in range, not just the picked one.
        option_ids = {i.pk for i in resp.context['ilap_options']}
        assert option_ids == {
            tikets[k].id_periode_data.id_sub_jenis_data_ilap.id_ilap_id
            for k in ('adhoc_recent', 'adhoc_lowercase', 'adhoc_refinished')
        }


@pytest.mark.django_db
class TestGenerate:
    def _post_data(self, *ids):
        return {
            'tanggal_mulai': (date.today() - timedelta(days=30)).isoformat(),
            'tanggal_selesai': date.today().isoformat(),
            'ilap_id': '',
            'ticket_ids': [str(i) for i in ids],
        }

    def test_generates_docx_from_pdi_template(self, client, pmde_user, tikets):
        _nd_template('ND PENGANTAR KE PDI')
        client.force_login(pmde_user)
        resp = client.post(reverse(URL_NAME), self._post_data(tikets['adhoc_recent'].pk))

        assert resp.status_code == 200
        assert resp['Content-Type'] == DOCX_MIME
        assert 'nd_pengantar_pdi_' in resp['Content-Disposition']
        doc = Document(BytesIO(resp.content))
        assert 'ND PENGANTAR KE PDI' in [p.text for p in doc.paragraphs]
        rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
        assert rows == [[tikets['adhoc_recent'].nomor_tiket, 'KPDE_ADHOC_BI_PADAN']]

    def test_ids_outside_the_listing_are_ignored(self, client, pmde_user, tikets):
        _nd_template('ND PENGANTAR KE PDI')
        client.force_login(pmde_user)
        resp = client.post(reverse(URL_NAME), self._post_data(
            tikets['adhoc_recent'].pk, tikets['regular_recent'].pk, tikets['adhoc_old'].pk,
        ))
        doc = Document(BytesIO(resp.content))
        assert [r.cells[0].text for r in doc.tables[0].rows] == [tikets['adhoc_recent'].nomor_tiket]

    def test_nothing_selected_redirects_back_with_filter(self, client, pmde_user, tikets):
        client.force_login(pmde_user)
        resp = client.post(reverse(URL_NAME), self._post_data(tikets['regular_recent'].pk))

        assert resp.status_code == 302
        assert resp.url.startswith(reverse(URL_NAME) + '?')
        assert 'tanggal_mulai=' in resp.url
